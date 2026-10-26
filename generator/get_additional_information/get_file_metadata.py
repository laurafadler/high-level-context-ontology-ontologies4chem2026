"""
Collect the metadata of all files in a folder (and its subfolders) and write it to an
Excel overview.

Used in two ways:
- by the crate generator (creation/media_object_functions.py), which reads the metadata of
  the series and experiment folders fresh on every run (collect_folder_metadata);
- on its own, to look at the metadata of a folder:
      python get_additional_information/get_file_metadata.py
      (then enter the folder path)

The overview lands in the folder itself as "<folder name>_collected_file_metadata.xlsx".
Overviews written earlier, and Office lock files ("~$..."), are never counted as files.

The header row of a table is always the FIRST row of the first sheet (or the first line of
a CSV file). A file whose first row is empty, or holds a number, a date or a time, has no
headers and counts as not tabular.

Important: all the files in the folder need to be closed for a complete list.
"""

# Needs openpyxl; pypdf only for the author of PDF files (if it is missing, the author stays empty):
# python -m pip install openpyxl pypdf

import sys
from pathlib import Path
from datetime import datetime, date, time
import csv
import io
import json
import zipfile
import xml.etree.ElementTree as ElementTree

import openpyxl

# Runs from generator/<folder>/ - the generator has to be on the search path.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from helpers.safe_write import write_atomic  # noqa: E402


OUTPUT_SUFFIX = "_collected_file_metadata.xlsx"

# Separator for the header line in the Excel overview - like the lists in the answer file
# (helpers/identifiers.py: LIST_SEPARATOR); a comma can stand in a header.
HEADER_SEPARATOR = "; "

CSV_ENCODINGS = ["utf-8-sig", "utf-8", "cp1252", "latin-1"]

COLUMNS = [
    "File Name",
    "Folder",
    "File Path",
    "File Type",
    "Size (bytes)",
    "Created",
    "Last Modified",
    "Author",
    "Character Encoding",
    "Headers",
]


# Word, Excel and PowerPoint are zip archives; the author stands in docProps/core.xml. So
# neither python-docx nor python-pptx is needed for it.
OFFICE_EXTENSIONS = {".docx", ".docm", ".xlsx", ".xlsm", ".pptx", ".pptm"}
DC_CREATOR = "{http://purl.org/dc/elements/1.1/}creator"


def get_author(file):
    """Try to extract the author from Office and PDF files; "" if there is none."""

    extension = file.suffix.lower()

    try:
        if extension in OFFICE_EXTENSIONS:
            with zipfile.ZipFile(file) as archive:
                core = ElementTree.fromstring(archive.read("docProps/core.xml"))
            creator = core.find(DC_CREATOR)
            return (creator.text or "").strip() if creator is not None else ""

        if extension == ".pdf":
            from pypdf import PdfReader
            metadata = PdfReader(file).metadata
            return (metadata.author or "") if metadata else ""

    except Exception:
        return ""

    return ""


def get_csv_encoding(file):
    """Try to determine the encoding of a CSV file."""

    if file.suffix.lower() != ".csv":
        return ""

    for encoding in CSV_ENCODINGS:
        try:
            with open(file, "r", encoding=encoding) as csv_file:
                csv_file.read(4096)
            return encoding
        except UnicodeDecodeError:
            continue
        except Exception:
            return ""

    return ""


def _is_data_cell(value):
    """Whether a cell of the first row holds data rather than a header: a number, a date or
    a time. Excel says so through the cell type; a CSV cell is text and counts as data when
    it reads as a number (decimal point or comma) or as an ISO date/time."""
    if isinstance(value, (bool, int, float, datetime, date, time)):
        return True
    text = str(value).strip()
    try:
        float(text.replace(",", "."))
        return True
    except ValueError:
        pass
    try:
        datetime.fromisoformat(text)
        return True
    except ValueError:
        return False


def get_headers(file):
    """The column headers of an Excel or CSV file: the non-empty cells of the first row,
    in order. Empty list for other file types, for a file whose first row is empty, and for a
    file whose first row holds a number, a date or a time - that row is data, the file has no
    header row and is not treated as a table (a note says so)."""

    extension = file.suffix.lower()

    try:
        # Excel
        if extension in [".xlsx", ".xlsm"]:
            workbook = openpyxl.load_workbook(file, read_only=True, data_only=True)
            worksheet = workbook.worksheets[0]
            first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), [])
            workbook.close()
            cells = [value for value in first_row if value is not None and str(value).strip()]

        # CSV
        elif extension == ".csv":
            encoding = get_csv_encoding(file)
            if not encoding:
                return []
            with open(file, "r", encoding=encoding, newline="") as csv_file:
                sample = csv_file.read(4096)
                csv_file.seek(0)
                # Recognize the separator: a European Excel writes CSV with semicolons, and
                # "time;T" would otherwise be a single header.
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
                except csv.Error:
                    dialect = csv.excel
                cells = [cell for cell in next(csv.reader(csv_file, dialect), []) if cell.strip()]

        else:
            return []

    except Exception:
        return []

    # A first line with numbers or dates is data, not a header. The file then counts as not
    # tabular, instead of getting numbers as column names. No stop as with duplicate headers:
    # raw data of a device must not be changed afterwards.
    if any(_is_data_cell(cell) for cell in cells):
        print(f"Note: the first row of {file.name} holds numbers or dates - it is read as data without a "
              "header row, so the file gets no columns.")
        return []
    return [str(cell).strip() for cell in cells]


def _eln_export_date(file):
    """The moment an ELN export (.eln, The ELN Consortium) was written, from its own ro-crate-metadata.json:
    dateCreated of the metadata descriptor, else datePublished of its root - an ISO string as the ELN gives it
    (without a time zone if it has none). None if the file says nothing. Since 04.10.2026: the date in the
    file system is the moment the export was copied into the input folder, not the export (Laura Fadler)."""
    try:
        with zipfile.ZipFile(file) as archive:
            name = next(n for n in archive.namelist() if n.endswith("ro-crate-metadata.json"))
            graph = {e.get("@id"): e for e in json.loads(archive.read(name)).get("@graph", [])}
    except (zipfile.BadZipFile, StopIteration, ValueError, OSError):
        return None
    descriptor = graph.get("ro-crate-metadata.json", {})
    root = graph.get((descriptor.get("about") or {}).get("@id", "./"), {})
    value = descriptor.get("dateCreated") or root.get("datePublished")
    return str(value).strip().replace(" ", "T", 1) if value else None


def get_file_metadata(file, root):
    """Collect metadata for one file. root is the folder the scan started from; "Folder"
    is the file's subfolder relative to it ("" directly in root)."""

    file_stats = file.stat()
    # Note: on Windows st_ctime is the date of creation of the file AT THIS PLACE - after copying
    # (not moving) that can be the date of the copy, not that of the measurement.
    created = datetime.fromtimestamp(file_stats.st_ctime).astimezone()
    modified = datetime.fromtimestamp(file_stats.st_mtime).astimezone()
    relative_folder = file.parent.relative_to(root).as_posix()

    return {
        "File Name": file.name,
        "Folder": "" if relative_folder == "." else relative_folder,
        "File Path": str(file.resolve()),
        "File Type": file.suffix.lower().lstrip("."),
        "Size (bytes)": file_stats.st_size,
        "Created": (_eln_export_date(file) if file.suffix.lower() == ".eln" else None) or created.isoformat(timespec="seconds"),
        "Last Modified": modified.isoformat(timespec="seconds"),
        "Author": get_author(file),
        "Character Encoding": get_csv_encoding(file),
        "Headers": get_headers(file),
    }


def _is_collected_file(file):
    """Whether a file belongs to the folder's content: not an overview written by this
    script, not an Office lock file, not hidden."""
    return not (file.name.endswith(OUTPUT_SUFFIX) or file.name.startswith("~$") or file.name.startswith("."))


def collect_folder_metadata(folder):
    """Metadata of every file in folder and its subfolders, sorted by folder and file name.

    Returns (rows, failed): rows as from get_file_metadata, failed as (file name, reason) for
    files that could not be read.
    """
    folder = Path(folder)
    rows, failed = [], []
    for file in sorted(folder.rglob("*"), key=lambda f: (f.parent.relative_to(folder).as_posix(), f.name)):
        if not file.is_file() or not _is_collected_file(file):
            continue
        try:
            rows.append(get_file_metadata(file, folder))
        except Exception as error:
            failed.append((file.name, str(error)))
    return rows, failed


def overview_path(folder):
    """Where the overview of folder is written inside folder itself."""
    folder = Path(folder)
    return folder / f"{folder.name}{OUTPUT_SUFFIX}"


def write_metadata_workbook(rows, output_file):
    """Write rows (see collect_folder_metadata) as an Excel overview to output_file."""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "File Metadata"
    worksheet.append(COLUMNS)
    for row in rows:
        worksheet.append([HEADER_SEPARATOR.join(row[c]) if c == "Headers" else row[c] for c in COLUMNS])

    for letter, width in zip("ABCDEFGHIJ", [40, 25, 80, 12, 15, 30, 30, 30, 20, 60]):
        worksheet.column_dimensions[letter].width = width

    # Into memory first, then atomically onto the disk (helpers/safe_write.py): the overview lies
    # next to the crate and in the data folder, both in OneDrive.
    buffer = io.BytesIO()
    workbook.save(buffer)
    write_atomic(output_file, buffer.getvalue())


def main():
    folder = Path(input("Enter folder path: ").strip().strip('"'))
    if not folder.is_dir():
        print("The specified folder does not exist.")
        return

    rows, failed = collect_folder_metadata(folder)
    output_file = overview_path(folder)
    write_metadata_workbook(rows, output_file)

    print("\nDone!")
    print(f"Metadata of {len(rows)} file(s) saved to:\n{output_file}")
    if failed:
        print("\nThe following files could not be processed:")
        for filename, reason in failed:
            print(f"  - {filename}: {reason}")
    else:
        print("\nAll files were processed successfully.")


if __name__ == "__main__":
    main()
